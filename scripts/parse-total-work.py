"""Turn Total work.xlsx into roster JSON for scripts/import-roster.mjs."""
import json
import re
import sys

import openpyxl

PATH = r"C:\Users\javaz\Downloads\Total work.xlsx"

SHEETS = {
    "Eng14": ("SAT 14", "ebrw", {"name": 1, "phone": 4, "parent": 5, "grade": 6, "english": 7, "math": 8, "goal": 9}),
    "Math14": ("SAT 14", "math", {"name": 1, "phone": 3, "parent": 4, "grade": 5, "english": 6, "math": 7, "goal": 8}),
    "Eng13": ("SAT 13", "ebrw", {"name": 1, "phone": 3, "parent": 4, "grade": 5, "english": 7, "math": 8, "goal": 9}),
    "Math13": ("SAT 13", "math", {"name": 1, "phone": 3, "parent": 4, "grade": 5, "english": 6, "math": 7, "goal": 8}),
    "Eng12": ("SAT 12", "ebrw", {"name": 1, "phone": 3, "parent": 4, "grade": 5, "english": 6, "math": 7, "goal": 8}),
    "Math12": ("SAT 12", "math", {"name": 1, "phone": 2, "parent": 3, "grade": 4, "english": 5, "math": 6, "goal": 7}),
    "Eng10": ("SAT 10", "ebrw", {"name": 1, "phone": 3, "parent": 4, "grade": 5, "english": 6, "math": 7, "goal": 8}),
    "Math10": ("SAT 10", "math", {"name": 1, "phone": 3, "parent": 4, "grade": 5, "english": 6, "math": 7, "goal": 8}),
    "PreE": ("Pre-SAT", "ebrw", {"name": 1, "phone": 3, "parent": 4, "grade": 5, "english": 6, "math": 7, "goal": 8}),
    "PreM": ("Pre-SAT", "math", {"name": 1, "phone": 3, "parent": 4, "grade": 5, "english": 6, "math": 7, "goal": 8}),
}

RANK = {"SAT 14": 5, "SAT 13": 4, "SAT 12": 3, "SAT 10": 2, "Pre-SAT": 1}


def cell(ws, row, col):
    value = ws.cell(row, col).value
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def digits(text):
    return len(re.findall(r"\d", text or ""))


def as_phone(value):
    if not value:
        return None
    count = digits(value)
    if count < 9 or count > 15:
        return None
    if re.search(r"[A-Za-z]", value):
        return None
    return value


def as_grade(value):
    if not value:
        return None
    if re.search(r"sinf|bitir", value, re.I):
        return value
    if re.fullmatch(r"\d+(\.0)?", value):
        number = int(float(value))
        if 8 <= number <= 12:
            return str(number)
    return None


def as_note(value):
    if not value:
        return None
    if re.fullmatch(r"[\d.\s~+\-/]+", value):
        return None
    if re.fullmatch(r"(d|done|n\.?a\.?|-|\?\?\?|shu hafta|ertaga|continue|qilindi)", value, re.I):
        return None
    if len(value) < 2:
        return None
    return value[:200]


def norm_name(raw):
    text = raw.strip().lower().replace("'", "").replace("’", "")
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"\s+\d*\s*hafta.*$", "", text).strip()
    return text


def display_name(raw):
    text = re.sub(r"\s+", " ", raw.strip())
    text = re.sub(r"\s+\d*\s*hafta.*$", "", text, flags=re.I).strip()
    return text


def main():
    wb = openpyxl.load_workbook(PATH, read_only=True, data_only=True)
    by_class = {}
    for sheet, (class_name, subject, cols) in SHEETS.items():
        ws = wb[sheet]
        first = True
        for row in ws.iter_rows(min_row=1, max_col=12):
            if first:
                first = False
                continue
            values = {key: None for key in cols}
            raw_name = None
            for key, col in cols.items():
                if col - 1 >= len(row):
                    continue
                value = row[col - 1].value
                if value is None:
                    continue
                text = str(value).strip()
                if not text:
                    continue
                if key == "name":
                    raw_name = text
                else:
                    values[key] = text
            if not raw_name or raw_name in ("Name", "`"):
                continue
            if re.fullmatch(r"[\d\s+]+", raw_name):
                continue
            key = norm_name(raw_name)
            if not key:
                continue
            bucket = by_class.setdefault(class_name, {})
            person = bucket.get(key)
            if person is None:
                person = {
                    "className": class_name,
                    "name": display_name(raw_name),
                    "norm": key,
                    "tokens": len(key.split()),
                    "subjects": [],
                    "phone": None,
                    "parentPhone": None,
                    "grade": None,
                    "englishNote": None,
                    "mathNote": None,
                    "goal": None,
                }
                bucket[key] = person
            if subject not in person["subjects"]:
                person["subjects"].append(subject)
            person["phone"] = person["phone"] or as_phone(values["phone"])
            person["parentPhone"] = person["parentPhone"] or as_phone(values["parent"])
            person["grade"] = person["grade"] or as_grade(values["grade"])
            person["englishNote"] = person["englishNote"] or as_note(values["english"])
            person["mathNote"] = person["mathNote"] or as_note(values["math"])
            person["goal"] = person["goal"] or as_note(values["goal"])

    kept = {}
    for class_name, people in by_class.items():
        for key, person in people.items():
            if person["tokens"] < 2:
                kept[f"{class_name}|{key}"] = person
                continue
            current = kept.get(key)
            if current is None or RANK[class_name] > RANK[current["className"]]:
                kept[key] = person

    out = []
    for person in kept.values():
        person["importKey"] = f"{person['className']}|{person['norm']}"
        person.pop("tokens")
        out.append(person)
    out.sort(key=lambda row: (row["className"], row["name"].lower()))
    json.dump(out, sys.stdout, ensure_ascii=False)


if __name__ == "__main__":
    main()
